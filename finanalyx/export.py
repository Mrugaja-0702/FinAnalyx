"""Machine-readable exports: JSON and Excel."""
from __future__ import annotations

import json
import math
from dataclasses import asdict
from pathlib import Path

import numpy as np
import pandas as pd

from .analyzer import Analysis, health_scores
from .scores import compute_scores
from .model import EXTRA_LABELS, fy, item_label
from .ratios import HEADLINE, METRIC_INDEX, METRICS, fmt_value
from .schema import ITEMS

_ITEM_ORDER = {k: i for i, k in enumerate([*ITEMS, *EXTRA_LABELS])}


def to_dict(a: Analysis) -> dict:
    fd, rr, tr = a.data, a.ratios, a.trends
    return {
        "company": fd.company,
        "years": fd.years,
        "units": fd.units,
        "balance_basis": rr.basis,
        "year_labels": {str(y): fy(y) for y in fd.years},
        "health_score": health_scores(a),
        "scores": compute_scores(a.data),
        "line_items": {k: {"label": item_label(k),
                           "statement": ITEMS[k].statement if k in ITEMS else "derived",
                           "order": _ITEM_ORDER.get(k, 999),
                           "values": {str(y): v for y, v in vals.items()},
                           "source": {str(y): s for y, s in fd.provenance.get(k, {}).items()}}
                       for k, vals in fd.values.items() if vals},
        "headline": list(HEADLINE),
        "ratios": {
            m.key: {
                "name": m.name, "category": m.category, "unit": m.unit, "formula": m.formula,
                "description": m.description, "higher_is_better": m.higher_is_better,
                "good": m.good, "weak": m.weak, "period": m.period,
                "latest_year": a.ratios.latest(m.key)[0],
                **({"value": rr.get(m.key).value, "status": rr.get(m.key).status, "note": rr.get(m.key).note,
                    "health": m.health(rr.get(m.key))} if m.period else
                   {"by_year": {str(y): {"value": v.value, "status": v.status, "note": v.note, "health": m.health(v)}
                                for y, v in rr.per_year.get(m.key, {}).items()}}),
            } for m in METRICS
        },
        "dashboard": [
            {"metric": METRIC_INDEX[k].name, "year": y, "value": v.value, "status": v.status,
             "display": fmt_value(METRIC_INDEX[k], v), "health": METRIC_INDEX[k].health(v)}
            for k in HEADLINE for y, v in [a.headline_value(k)]
        ],
        "summary": tr.summary,
        "trends_note": tr.note,
        "insights": [asdict(i) for i in tr.insights],
        "year_over_year": [{"year": y, "text": t} for y, t in tr.yoy],
        "assumptions": fd.assumptions,
        "warnings": fd.warnings,
        "mapping": [asdict(r) for r in fd.mapping_log],
        "unmatched_labels": [{"table": t, "label": lbl} for t, lbl in fd.unmatched],
    }


def clean(obj):
    """Make a structure strictly JSON-safe (numpy scalars -> Python, NaN / infinity -> null)."""
    if isinstance(obj, np.generic):
        obj = obj.item()
    if isinstance(obj, float):
        return obj if math.isfinite(obj) else None
    if isinstance(obj, dict):
        return {str(k): clean(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [clean(v) for v in obj]
    return obj


def to_json(a: Analysis) -> str:
    return json.dumps(clean(to_dict(a)), indent=2, default=str)


def to_excel(a: Analysis, path: str | Path) -> None:
    fd, rr, tr = a.data, a.ratios, a.trends
    cols = [fy(y) for y in fd.years]

    dash = []
    for k in HEADLINE:
        m = METRIC_INDEX[k]
        y, v = a.headline_value(k)
        dash.append({"Metric": m.name, "Year": fy(y) if y else "period", "Value": v.value,
                     "Status": v.status if not v.ok else (m.health(v) or ""), "Note": v.note, "Formula": m.formula})

    ratio_rows = []
    for m in METRICS:
        row = {"Category": m.category, "Metric": m.name, "Unit": m.unit}
        if m.period:
            v = rr.get(m.key)
            row["Period"] = v.value
            row["Note"] = v.note
        else:
            for y, c in zip(fd.years, cols):
                v = rr.get(m.key, y)
                row[c] = v.value if v.ok else v.status
        row["Formula"] = m.formula
        ratio_rows.append(row)

    items = fd.to_frame()
    items.columns = cols
    insights = pd.DataFrame([{"Severity": i.severity, "Category": i.category, "Finding": i.title,
                              "Detail": i.detail, "Years": ", ".join(fy(y) for y in i.years), "Score": i.score}
                             for i in tr.insights])
    notes = pd.DataFrame([{"Type": "Summary", "Text": t} for t in tr.summary]
                         + [{"Type": "Assumption", "Text": t} for t in fd.assumptions]
                         + [{"Type": "Warning", "Text": t} for t in fd.warnings])

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(dash).to_excel(xw, sheet_name="Dashboard", index=False)
        pd.DataFrame(ratio_rows).to_excel(xw, sheet_name="Ratios", index=False)
        items.to_excel(xw, sheet_name="Line Items", index_label="Line item")
        insights.to_excel(xw, sheet_name="Trends", index=False)
        notes.to_excel(xw, sheet_name="Notes", index=False)
        _format_workbook(xw.book, ratio_rows, dash)


def _format_workbook(wb, ratio_rows: list[dict], dash: list[dict]) -> None:
    from openpyxl.styles import Font, PatternFill

    header_fill = PatternFill("solid", fgColor="EEEDE8")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        ws.freeze_panes = "B2" if ws.title in {"Ratios", "Line Items"} else "A2"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)

    fmt = {"%": "0.0%", "x": '0.00"x"', "days": '0" d"', "abs": "#,##0"}
    ws = wb["Ratios"]
    for r_idx, row in enumerate(ratio_rows, start=2):
        number_format = fmt[row["Unit"]]
        for cell in ws[r_idx][3:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format
    ws = wb["Dashboard"]
    for r_idx, row in enumerate(dash, start=2):
        unit = next(m.unit for m in METRICS if m.name == row["Metric"])
        ws.cell(r_idx, 3).number_format = fmt[unit]
    ws = wb["Line Items"]
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "#,##0.0;(#,##0.0)"
